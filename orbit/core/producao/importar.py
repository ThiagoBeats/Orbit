"""Leitura de planilhas (CSV/XLSX) com mapeamento de colunas (spec Produção §7).

CSV: separador `;`, `,` ou tab detectado sozinho; UTF-8 (com ou sem BOM) ou
Latin-1. XLSX: primeira planilha, via openpyxl. O cabeçalho é a primeira
linha não vazia; linhas vazias são ignoradas. Tudo vira texto.
"""
from __future__ import annotations

import csv
import io
from pathlib import Path

from orbit.core.errors import OrbitError

LIMITE_LINHAS = 20000


def _texto_celula(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _vazia(linha: list[str]) -> bool:
    return not any(c for c in linha)


def _ler_csv(caminho: Path) -> list[list[str]]:
    dados = caminho.read_bytes()
    for cod in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            texto = dados.decode(cod)
            break
        except UnicodeDecodeError:
            continue
    primeira = next((l for l in texto.splitlines() if l.strip()), "")
    sep = max((";", ",", "\t"), key=primeira.count)
    return [[c.strip() for c in l] for l in csv.reader(io.StringIO(texto), delimiter=sep)]


def _ler_xlsx(caminho: Path) -> list[list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover
        raise OrbitError("Leitura de .xlsx indisponível (openpyxl não instalado).") from None
    try:
        wb = load_workbook(caminho, read_only=True, data_only=True)
    except Exception as e:
        raise OrbitError(f"Não foi possível abrir a planilha: {e}") from e
    try:
        ws = wb.worksheets[0]
        out = []
        for row in ws.iter_rows(values_only=True):
            out.append([_texto_celula(v) for v in row])
            if len(out) > LIMITE_LINHAS:
                break
        return out
    finally:
        wb.close()


def ler_tabela(caminho) -> dict:
    caminho = Path(caminho)
    if not caminho.is_file():
        raise OrbitError("Arquivo não encontrado.")
    ext = caminho.suffix.lower()
    try:
        if ext in (".csv", ".txt"):
            linhas = _ler_csv(caminho)
        elif ext in (".xlsx", ".xlsm"):
            linhas = _ler_xlsx(caminho)
        else:
            raise OrbitError("Use uma planilha .xlsx ou .csv (salve o .xls como .xlsx).")
    except OSError as e:
        raise OrbitError(f"Não foi possível ler o arquivo: {e}") from e
    linhas = [l for l in linhas if not _vazia(l)]
    if not linhas:
        raise OrbitError("A planilha está vazia.")
    if len(linhas) > LIMITE_LINHAS:
        raise OrbitError(f"A planilha tem mais de {LIMITE_LINHAS} linhas.")
    cab = linhas[0]
    # remove colunas vazias do fim do cabeçalho (comum em xlsx)
    while cab and not cab[-1]:
        cab = cab[:-1]
    n = len(cab)
    corpo = [(l + [""] * n)[:n] for l in linhas[1:]]
    return {"cabecalho": cab, "linhas": [l for l in corpo if not _vazia(l)]}


def mapear(tabela: dict, mapa: dict, obrigatorios: list[str]) -> list[dict]:
    """Converte as linhas em dicts {campo: valor} pelo índice de coluna escolhido.

    Linhas sem os campos obrigatórios são ignoradas."""
    n = len(tabela.get("cabecalho") or [])
    for campo in obrigatorios:
        if mapa.get(campo) is None:
            raise OrbitError(f"Escolha a coluna de \"{campo}\".")
    for campo, idx in mapa.items():
        if idx is None:
            continue
        if isinstance(idx, bool) or not isinstance(idx, int) or not 0 <= idx < n:
            raise OrbitError(f"Coluna inválida para \"{campo}\".")
    out = []
    for linha in tabela.get("linhas") or []:
        item = {campo: (linha[idx] if idx is not None and idx < len(linha) else "")
                for campo, idx in mapa.items()}
        if all(item.get(c) for c in obrigatorios):
            out.append(item)
    return out
